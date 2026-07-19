from __future__ import annotations

import hashlib
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path, PurePath
from urllib.parse import urlparse

from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader
from sqlalchemy import select, update
from sqlalchemy.orm import Session

from app.models import DataValidationDecision, SourceDocument, UserAccount

MAX_SOURCE_BYTES = 100 * 1024 * 1024
DOCUMENT_TYPES = frozenset({"ADMISSION_GUIDE", "ADMISSION_RESULT", "NOTICE", "OTHER"})
ENTITY_TYPES = frozenset({"CATALOG", "RULE", "ADMISSION_RESULT"})
RESOLUTION_STATUSES = frozenset({"CONFIRMED", "REJECTED"})


class SourceDocumentError(ValueError):
    pass


def register_source_document(
    session: Session,
    *,
    storage_root: Path,
    filename: str,
    body: bytes,
    academic_year: str,
    document_type: str,
    institution_id: str,
    admission_round_id: str,
    original_url: str,
    announced_at: str,
    revision_label: str,
) -> SourceDocument:
    if not body or len(body) > MAX_SOURCE_BYTES:
        raise SourceDocumentError("출처 파일은 1바이트 이상 100MiB 이하여야 합니다.")
    safe_name = PurePath(filename).name.strip()
    if not safe_name or len(safe_name) > 255:
        raise SourceDocumentError("출처 파일 이름을 확인하세요.")
    suffix = Path(safe_name).suffix.lower()
    page_count = _validate_and_count(body, suffix)
    year = _required_year(academic_year)
    kind = document_type.strip().upper()
    if kind not in DOCUMENT_TYPES:
        raise SourceDocumentError("출처 문서 유형을 확인하세요.")
    source_url = _validated_url(original_url)
    announced = _optional_datetime(announced_at)
    revision = revision_label.strip()
    if len(revision) > 120:
        raise SourceDocumentError("개정 표시는 120자 이하여야 합니다.")
    digest = hashlib.sha256(body).hexdigest()
    if session.scalar(select(SourceDocument.id).where(SourceDocument.file_hash == digest)):
        raise SourceDocumentError("같은 파일이 이미 등록되어 있습니다.")
    normalized_institution_id = institution_id.strip() or None
    session.execute(
        update(SourceDocument)
        .where(
            SourceDocument.academic_year == year,
            SourceDocument.institution_id == normalized_institution_id,
            SourceDocument.document_type == kind,
            SourceDocument.is_current.is_(True),
        )
        .values(is_current=False)
    )
    record = SourceDocument(
        academic_year=year,
        institution_id=normalized_institution_id,
        admission_round_id=admission_round_id.strip() or None,
        document_type=kind,
        document_status="DRAFT",
        revision_label=revision or None,
        file_hash=digest,
        page_count=page_count,
        detected_years=[year],
        year_consistency_status="CONSISTENT",
        verification_status="PENDING",
        original_url=source_url,
        announced_at=announced,
        original_filename=safe_name,
        is_current=True,
    )
    session.add(record)
    session.flush()
    relative_path = Path("source-documents") / f"{record.id}{suffix}"
    target_root = storage_root.resolve()
    target = (target_root / relative_path).resolve()
    if target_root not in target.parents:
        raise SourceDocumentError("출처 파일 저장 경로가 유효하지 않습니다.")
    try:
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(body)
    except OSError as error:
        raise SourceDocumentError("출처 파일을 안전하게 저장할 수 없습니다.") from error
    record.storage_path = relative_path.as_posix()
    return record


def create_validation_decision(
    session: Session,
    *,
    source_document_id: str,
    entity_type: str,
    entity_reference: str,
    field_name: str,
    current_value: str,
    portal_value: str,
    document_value: str,
) -> DataValidationDecision:
    if session.get(SourceDocument, source_document_id) is None:
        raise SourceDocumentError("출처 문서를 찾을 수 없습니다.")
    normalized_type = entity_type.strip().upper()
    if normalized_type not in ENTITY_TYPES:
        raise SourceDocumentError("비교 대상 유형을 확인하세요.")
    reference = entity_reference.strip()
    field = field_name.strip()
    if not reference or len(reference) > 160 or not field or len(field) > 120:
        raise SourceDocumentError("비교 대상과 필드명을 확인하세요.")
    values = (current_value.strip(), portal_value.strip(), document_value.strip())
    if any(len(value) > 4000 for value in values):
        raise SourceDocumentError("비교 값은 각각 4,000자 이하여야 합니다.")
    record = DataValidationDecision(
        source_document_id=source_document_id,
        entity_type=normalized_type,
        entity_reference=reference,
        field_name=field,
        current_value=values[0] or None,
        portal_value=values[1] or None,
        document_value=values[2] or None,
        resolution_status="PENDING",
    )
    session.add(record)
    return record


def resolve_validation_decision(
    session: Session,
    *,
    decision_id: str,
    user: UserAccount,
    resolution_status: str,
    resolved_value: str,
    resolution_reason: str,
) -> DataValidationDecision:
    record = session.get(DataValidationDecision, decision_id)
    if record is None:
        raise SourceDocumentError("검증 항목을 찾을 수 없습니다.")
    status = resolution_status.strip().upper()
    value = resolved_value.strip()
    reason = resolution_reason.strip()
    if status not in RESOLUTION_STATUSES:
        raise SourceDocumentError("검증 결정을 확인하세요.")
    if not value or len(value) > 4000 or not reason or len(reason) > 2000:
        raise SourceDocumentError("확정 값과 결정 이유를 입력하세요.")
    record.resolution_status = status
    record.resolved_value = value
    record.resolution_reason = reason
    record.reviewed_by_user_account_id = user.id
    return record


def _validate_and_count(body: bytes, suffix: str) -> int:
    try:
        if suffix == ".pdf" and body.startswith(b"%PDF-"):
            return max(1, len(PdfReader(BytesIO(body)).pages))
        if suffix == ".png" and body.startswith(b"\x89PNG\r\n\x1a\n"):
            with Image.open(BytesIO(body)) as image:
                image.verify()
            return 1
        if suffix in {".jpg", ".jpeg"} and body.startswith(b"\xff\xd8\xff"):
            with Image.open(BytesIO(body)) as image:
                image.verify()
            return 1
        if suffix == ".xlsx" and body.startswith(b"PK"):
            workbook = load_workbook(BytesIO(body), read_only=True, data_only=True)
            page_count = max(1, len(workbook.worksheets))
            workbook.close()
            return page_count
        if suffix == ".csv":
            body.decode("utf-8-sig")
            return 1
    except (OSError, UnicodeError, ValueError) as error:
        raise SourceDocumentError("출처 파일 내용이 확장자와 일치하지 않습니다.") from error
    raise SourceDocumentError("PDF, PNG, JPG, CSV, XLSX 출처 파일만 등록할 수 있습니다.")


def _required_year(raw: str) -> int:
    try:
        year = int(raw.strip())
    except ValueError as error:
        raise SourceDocumentError("학년도를 확인하세요.") from error
    if not 2000 <= year <= 2100:
        raise SourceDocumentError("학년도는 2000 이상 2100 이하입니다.")
    return year


def _validated_url(raw: str) -> str | None:
    value = raw.strip()
    if not value:
        return None
    parsed = urlparse(value)
    if (
        len(value) > 1000
        or parsed.scheme not in {"http", "https"}
        or not parsed.hostname
        or parsed.username
        or parsed.password
    ):
        raise SourceDocumentError("원문 URL은 사용자 정보가 없는 HTTP(S) 주소여야 합니다.")
    return value


def _optional_datetime(raw: str) -> datetime | None:
    value = raw.strip()
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(value)
    except ValueError as error:
        raise SourceDocumentError("발표일을 확인하세요.") from error
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=UTC)
    return parsed


__all__ = [
    "DOCUMENT_TYPES",
    "ENTITY_TYPES",
    "SourceDocumentError",
    "create_validation_decision",
    "register_source_document",
    "resolve_validation_decision",
]
